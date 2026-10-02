#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
actualizar_pnd_ps.py
====================
Carga de las fichas SGAPPG (Matriz de Planificación / Matriz de Seguimiento)
del PND · Protección Social e Inclusión Económica hacia los CSV del portal.

Estructura que mantiene (datos/pnd/proteccion_social/):
    pnd_ps_indicadores.csv            <- global, se edita a mano (casi nunca cambia)
    <anio>/acciones.csv               <- catálogo del año (1 fila por acción ÚNICA)
    <anio>/indicador_accion.csv       <- qué acción aparece en qué indicador
    <anio>/metas_programadas.csv      <- meta de cada mes (para "cumplimiento a la fecha")
    <anio>/actores.csv                <- actores clave por acción
    <anio>/seguimiento.csv            <- 1 fila por acción x corte (crece cada mes)
    <anio>/justificaciones.csv        <- textos largos, se cargan solo al abrir el detalle

Uso
---
  # Una vez al año, cuando llega la matriz de planificación nueva (genera BORRADOR
  # del catálogo; revisar acciones.csv y metas_programadas.csv antes de seguir):
  python scripts/actualizar_pnd_ps.py inicializar --anio 2026 --fichas ruta/fichas_2026-07

  # Cada mes:
  python scripts/actualizar_pnd_ps.py actualizar --periodo 2026-08 --fichas ruta/fichas_2026-08

  # Rama PND · Trabajo y Oportunidades (mismas fichas del mes; otra carpeta de datos):
  python scripts/actualizar_pnd_ps.py actualizar --periodo 2026-08 --fichas ruta/fichas_2026-08 \
         --datos datos/pnd/trabajo --indicadores indicadores.csv --reportes reportes/pnd_trabajo

Opciones comunes:
  --datos   carpeta raíz de datos (por defecto: <repo>/datos/pnd/proteccion_social)
  --reportes carpeta de reportes de validación (por defecto: <repo>/reportes/pnd_ps)

El modo "actualizar" es idempotente: si se vuelve a correr un periodo, reemplaza las
filas de ese periodo (no duplica).
"""
import argparse
import csv
import datetime as dt
import difflib
import os
import re
import sys
import unicodedata
import warnings
from collections import OrderedDict, defaultdict

warnings.filterwarnings("ignore")  # openpyxl: "Data Validation extension is not supported"
try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("Falta openpyxl: pip install openpyxl")

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATOS_DEFAULT = os.path.join(REPO, "datos", "pnd", "proteccion_social")
REPORTES_DEFAULT = os.path.join(REPO, "reportes", "pnd_ps")

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]
MES_NUM = {m: i + 1 for i, m in enumerate(MESES)}
MES_NUM["setiembre"] = 9

# Tipos de meta: definen cómo se acumula la meta programada hasta la fecha
#   flujo       -> se suma mes a mes (capacitados, créditos, talleres...)
#   cobertura   -> stock: la meta de cada mes es la meta vigente (usuarios atendidos)
#   trimestral  -> metas en marzo/junio/septiembre/diciembre, se suman
#   semestral   -> metas en junio/diciembre, se suman
#   hito        -> producto único; meta = 100 % en el mes de fecha fin
#   anual       -> meta anual sin programación mensual (no hay cumplimiento a la fecha)
#   tasa        -> meta expresada en % para cada mes (p. ej. 97,32 % de requerimientos a tiempo)
TIPOS_META = {"flujo", "cobertura", "trimestral", "semestral", "hito", "anual", "tasa"}

# ----------------------------------------------------------------------------- utilidades


def norm(s):
    s = "" if s is None else str(s)
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode()
    s = re.sub(r"[^a-z0-9 ]+", " ", s.lower())
    return re.sub(r"\s+", " ", s).strip()


def parse_num(v):
    """Número desde celda: acepta 14453, '14.453 ciudadanos', '265,000', '0,7', '97,32%'."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s or s.startswith("#") or norm(s) in ("no aplica", "n a", "na"):
        return None
    m = re.search(r"-?\d[\d.,]*", s)
    if not m:
        return None
    t = m.group(0).rstrip(".,")
    if re.fullmatch(r"\d{1,3}(\.\d{3})+", t):          # 14.453  -> miles con punto
        t = t.replace(".", "")
    elif re.fullmatch(r"\d{1,3}(,\d{3})+", t):         # 265,000 -> miles con coma
        t = t.replace(",", "")
    elif "," in t and "." in t:                         # 1.234,56
        t = t.replace(".", "").replace(",", ".")
    elif "," in t:                                      # 0,7
        t = t.replace(",", ".")
    try:
        return float(t)
    except ValueError:
        return None


def fnum(x, dec=2):
    if x is None or x == "":
        return ""
    x = round(float(x), dec)
    return str(int(x)) if x == int(x) else str(x)


def fecha_iso(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    s = str(v or "").strip()
    m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    return s[:10]


def slug(texto, usados, max_palabras=5):
    stop = {"de", "del", "la", "las", "los", "el", "en", "a", "y", "para", "por", "con",
            "que", "sus", "su", "al", "se", "o", "un", "una", "e", "como", "mtdh", "través"}
    palabras = [p for p in norm(texto).split() if p not in stop][:max_palabras]
    base = "-".join(palabras) or "accion"
    s, i = base, 2
    while s in usados:
        s, i = f"{base}-{i}", i + 1
    usados.add(s)
    return s


def leer_csv(ruta):
    if not os.path.exists(ruta):
        return []
    with open(ruta, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def escribir_csv(ruta, filas, campos):
    os.makedirs(os.path.dirname(ruta), exist_ok=True)
    with open(ruta, "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=campos, extrasaction="ignore")
        w.writeheader()
        for f in filas:
            w.writerow(f)


def enlaces(texto):
    return [u.rstrip(").,;") for u in re.findall(r"https?://\S+", str(texto or ""))]


# ----------------------------------------------------------------------------- lectura de fichas


def _secciones(ws):
    sec = {}
    for r in range(1, ws.max_row + 1):
        a = str(ws.cell(r, 1).value or "")
        m = re.match(r"SECCI[OÓ]N\s+(\d)", a)
        if m:
            sec[int(m.group(1))] = r
    return sec


def _filas_numeradas(ws, desde, hasta):
    """Filas entre dos secciones cuya columna A es un número y la B un texto válido."""
    out = []
    for r in range(desde + 1, hasta):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if parse_num(a) is None or b is None or str(b).startswith("#"):
            continue
        if not isinstance(a, (int, float)) and not re.fullmatch(r"\s*\d+\s*", str(a)):
            continue
        out.append(r)
    return out


def leer_ficha(ruta):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    P, S = wb["Matriz de Planificación"], wb["Matriz de Seguimiento"]
    info = {
        "eje": P["D8"].value, "objetivo": P["D9"].value, "meta": P["D10"].value,
        "indicador": P["D11"].value, "entidad": P["D14"].value,
    }
    sp, ss = _secciones(P), _secciones(S)
    plan = []
    for r in _filas_numeradas(P, sp[3], sp[4]):
        g = lambda c: P[f"{c}{r}"].value
        plan.append({
            "num": int(parse_num(g("A"))), "titulo": str(g("B")).strip(), "descripcion": str(g("C") or "").strip(),
            "tipo": g("D"), "financiamiento": g("E"), "fuente": g("F"), "cooperante": g("G"),
            "frecuencia": g("H"), "cobertura": g("I"), "territorio": g("J"),
            "fecha_inicio": fecha_iso(g("K")), "fecha_fin": fecha_iso(g("L")),
            "corriente": parse_num(g("M")), "inversion": parse_num(g("N")),
        })
    seg = []
    for r in _filas_numeradas(S, ss[3], ss[4]):
        g = lambda c: S[f"{c}{r}"].value
        seg.append({
            "num": int(parse_num(g("A"))), "titulo": str(g("B")).strip(), "descripcion": str(g("C") or "").strip(),
            "resultado_raw": g("K"), "avance_raw": g("L"), "justificacion": str(g("M") or "").strip(),
            "fecha_inicio": fecha_iso(g("N")), "fecha_fin": fecha_iso(g("O")), "estado": (str(g("P")).strip() if g("P") else ""),
            "corr_codificado": parse_num(g("Q")), "corr_reprogramado": parse_num(g("R")), "corr_devengado": parse_num(g("S")),
            "inv_codificado": parse_num(g("U")), "inv_reprogramado": parse_num(g("V")), "inv_devengado": parse_num(g("W")),
        })
    actores = []
    for r in _filas_numeradas(P, sp[4], sp[5]):
        g = lambda c: P[f"{c}{r}"].value
        actores.append({"num": int(parse_num(g("A"))), "actor": str(g("D") or "").strip(), "rol": str(g("F") or "").strip(),
                        "descripcion_rol": str(g("H") or "").strip(), "incidencia": str(g("L") or "").strip()})
    dificultades = {}
    if 5 in ss and 6 in ss:
        for r in _filas_numeradas(S, ss[5], ss[6]):
            txt = str(S[f"M{r}"].value or "").strip()
            if txt:
                dificultades[int(parse_num(S[f"A{r}"].value))] = txt
    return {"archivo": os.path.basename(ruta), "info": info, "plan": plan, "seg": seg,
            "actores": actores, "dificultades": dificultades}


def leer_fichas(carpeta, indicadores):
    """Devuelve {indicador_id: ficha} y la lista de archivos no reconocidos."""
    por_nombre = {norm(i["nombre"]): i["indicador_id"] for i in indicadores}
    fichas, ignoradas = {}, []
    for nombre in sorted(os.listdir(carpeta)):
        if not nombre.lower().endswith((".xlsx", ".xlsm")) or nombre.startswith("~$"):
            continue
        try:
            f = leer_ficha(os.path.join(carpeta, nombre))
        except Exception as e:  # ficha con otra plantilla
            ignoradas.append((nombre, f"no se pudo leer ({e.__class__.__name__}: {e})"))
            continue
        iid = por_nombre.get(norm(f["info"]["indicador"]))
        if not iid:
            ignoradas.append((nombre, f"indicador fuera de Protección Social: {f['info']['indicador']}"))
            continue
        if iid in fichas:
            ignoradas.append((nombre, f"ficha duplicada para {iid} (se usa {fichas[iid]['archivo']})"))
            continue
        fichas[iid] = f
    return fichas, ignoradas


# ----------------------------------------------------------------------------- metas desde la descripción


def _mes_fin(fecha_fin, anio):
    try:
        y, m = int(fecha_fin[:4]), int(fecha_fin[5:7])
        return m if y == anio else 12
    except Exception:
        return 12


def parse_metas(desc, anio, fecha_fin):
    """Borrador de (tipo_meta, meta_anual, {mes: meta_del_mes}, nota) a partir del texto.
    Es un BORRADOR: se revisa a mano una vez al año en metas_programadas.csv."""
    d = " ".join(str(desc or "").split())
    dn = norm(d)
    # 1) lista mensual explícita: "Enero: 978 Febrero: 978 ..."
    pares = re.findall(r"\b(enero|febrero|marzo|abril|mayo|junio|julio|agosto|septiembre|setiembre|octubre|noviembre|diciembre)\s*:?\s*(\d[\d.,]*)\s*(%?)",
                       d, flags=re.I)
    meses, meses_pct = OrderedDict(), OrderedDict()
    for mes, val, pct in pares:
        mnum = MES_NUM[mes.lower()]
        destino = meses_pct if pct else meses
        if mnum not in destino:
            destino[mnum] = parse_num(val)
    # 1a) metas en porcentaje mes a mes ("Enero: 99,04% ... Total 2026: 97,32%") -> tasa
    if len(meses_pct) >= 10:
        m_total = re.search(rf"[Tt]otal\s*{anio}\s*:\s*(\d[\d.,]*)\s*%", d)
        total = parse_num(m_total.group(1)) if m_total else meses_pct[max(meses_pct)]
        return "tasa", total, dict(meses_pct), "Meta en porcentaje para cada mes (tasa)."
    # 1b) tramos "enero-febrero: 13, marzo-abril: 17, Julio: 22 ..." (la meta se asigna al último mes del tramo)
    nombres = "|".join(MES_NUM)
    tramos = re.findall(rf"\b({nombres})\s*(?:-\s*({nombres}))?\s*:\s*(\d[\d.,]*)", d, flags=re.I)
    if any(t[1] for t in tramos):
        meses = OrderedDict()
        for m1, m2, val in tramos:
            meses[MES_NUM[(m2 or m1).lower()]] = parse_num(val)
        m_an = re.search(r"meta anual(?: referencial)?\s*(?:es|de|son)?\s*(\d[\d.,]*)", d, re.I)
        total = parse_num(m_an.group(1)) if m_an else sum(v for v in meses.values() if v)
        return "flujo", total, dict(meses), "Programación por tramos (bimestres/meses) tomada de la descripción."
    if len(meses) >= 10:
        m_total = re.search(rf"{anio}\s*:\s*(\d[\d.,]*)", d)
        total = parse_num(m_total.group(1)) if m_total else sum(v for v in meses.values() if v)
        return "flujo", total, dict(meses), "Programación mensual tomada de la descripción."
    # 2) trimestral: "IT. 10% (1.332) IIT: 6,67% (888) ..."
    trims = re.findall(r"\b(IV|III|II|I)T[.:]?\s*[\d.,]+\s*%\s*\((\d[\d.,]*)\)", d)
    if len(trims) >= 3:
        mes_de = {"I": 3, "II": 6, "III": 9, "IV": 12}
        metas = {mes_de[t]: parse_num(v) for t, v in trims}
        m_an = re.search(r"\((\d[\d.,]*)\s*acciones\)", d) or re.search(r"meta anual (?:son|es|de)?\s*(\d[\d.,]*)", d, re.I)
        total = parse_num(m_an.group(1)) if m_an else sum(metas.values())
        return "trimestral", total, metas, "Metas trimestrales tomadas de la descripción (acumuladas)."
    # 3) semestral con meta del primer semestre
    m_s1 = re.search(r"primer semestre es de\s*(\d[\d.,]*)", d, re.I)
    if "semestral" in dn and m_s1:
        s1 = parse_num(m_s1.group(1))
        m_an = re.search(r"un total de\s*(\d[\d.,]*)", d, re.I)
        total = parse_num(m_an.group(1)) if m_an else None
        metas = {6: s1}
        if total:
            metas[12] = total - s1
        return "semestral", total, metas, "Reporte semestral: meta del 1er semestre y total anual (medidas distintas; revisar)."
    # 4) cobertura (stock mensual)
    if "mensual es equivalente a la meta anual" in dn or "meta mensual es equivalente" in dn or "frecuencia mensual" in dn:
        m_c = (re.search(r"cobertura de\s*(\d[\d.,]*)", d, re.I) or re.search(r"es de\s*(\d[\d.,]*)", d, re.I)
               or re.search(r"[Aa]tenci[oó]n mensual a\s*(\d[\d.,]*)", d) or re.search(r"aproximadamente\s*(\d[\d.,]*)", d))
        if m_c:
            v = parse_num(m_c.group(1))
            return "cobertura", v, {m: v for m in range(1, 13)}, "Cobertura: la meta de cada mes es la meta anual."
    # 5) flujo con meta anual + mensual aproximada
    m_an = (re.search(r"[Mm]eta [Aa]nual\s*(?:de|son|es de|referencial)?\s*(\d[\d.,]*)", d)
            or re.search(r"planificado la entrega de\s*(\d[\d.,]*)", d)
            or re.search(r"meta de\s*(\d[\d.,]*)\s*participaciones", d))
    m_me = (re.search(r"[Mm]eta [Mm]ensual\s*(\d[\d.,]*)", d)
            or re.search(r"(\d[\d.,]*)\s*(?:usuarios|participaciones|capacitaciones)?\s*(?:de manera )?mensual", d))
    if m_an and m_me:
        tot, me = parse_num(m_an.group(1)), parse_num(m_me.group(1))
        if tot and me and me < tot:
            metas = {m: me for m in range(1, 12)}
            metas[12] = round(tot - me * 11, 2)
            return "flujo", tot, metas, "Meta mensual aproximada (dic. ajusta al total anual)."
    # 6) hito / producto único
    if "100%" in d.replace(" ", "") or "hoja de ruta" in dn or "reporta de manera anual" in dn:
        return "hito", 100.0, {_mes_fin(fecha_fin, anio): 100.0}, "Producto/hito: meta 100 % a la fecha de fin."
    # 7) meta anual sin programación
    m_ses = re.search(r"(\d+)\s*sesiones", d)
    if m_ses:
        v = parse_num(m_ses.group(1))
        return "anual", v, {12: v}, "Meta anual sin programación mensual."
    if not m_an:
        m_an = re.search(r"meta anual\b[^.]{0,80}?\bes de\s*(\d[\d.,]*)", d, re.I)
    if m_an:
        v = parse_num(m_an.group(1))
        return "anual", v, {12: v}, "Meta anual sin programación mensual."
    return "anual", None, {}, "SIN META IDENTIFICABLE: completar a mano."


def meta_a_la_fecha(tipo, metas_mes, mes):
    """metas_mes: {mes: meta_del_mes}. Devuelve la meta acumulada/vigente al mes dado."""
    if not metas_mes:
        return None
    if tipo in ("cobertura", "tasa"):
        vals = [metas_mes[m] for m in sorted(metas_mes) if m <= mes]
        return vals[-1] if vals else None
    if tipo == "hito":
        return max(metas_mes.values()) if any(m <= mes for m in metas_mes) else None
    if tipo == "anual":
        return None
    acum = sum(v for m, v in metas_mes.items() if m <= mes and v is not None)
    return acum if acum > 0 else None


# ----------------------------------------------------------------------------- INICIALIZAR (anual)


def cmd_inicializar(args):
    anio = int(args.anio)
    dir_anio = os.path.join(args.datos, str(anio))
    destino = os.path.join(dir_anio, "acciones.csv")
    if os.path.exists(destino) and not args.forzar:
        sys.exit(f"Ya existe {destino}. Use --forzar para regenerar el borrador (se sobrescribe el catálogo).")
    indicadores = sorted(leer_csv(os.path.join(args.datos, args.indicadores)), key=lambda r: int(r["orden"]))
    fichas, ignoradas = leer_fichas(args.fichas, indicadores)

    catalogo, puente, metas, actores, usados = [], [], [], [], set()
    for ind in indicadores:
        iid = ind["indicador_id"]
        f = fichas.get(iid)
        if not f:
            print(f"  ! Sin ficha para {iid}")
            continue
        for a in f["plan"]:
            clave = norm(a["titulo"])
            _t, _anual, _m, _n = parse_metas(a["descripcion"], anio, a["fecha_fin"])
            match = None
            for c in catalogo:
                r = difflib.SequenceMatcher(None, clave, norm(c["titulo"])).ratio()
                # misma acción: título casi idéntico, o muy parecido Y con la misma meta anual
                # (evita fusionar p. ej. "Capacitar a los ciudadanos…" con "Capacitar a jóvenes…")
                rd = difflib.SequenceMatcher(None, norm(a["descripcion"]), norm(c["descripcion"])).ratio()
                if r >= 0.95 or (r >= 0.85 and rd >= 0.9 and _anual is not None and fnum(_anual) == c["meta_anual"]):
                    match = c
                    break
            if match is None:
                tipo_meta, anual, metas_mes, nota = parse_metas(a["descripcion"], anio, a["fecha_fin"])
                comp = ""
                dn = norm(a["descripcion"])
                if "inversion" in dn and "corriente" in dn and re.search(r"Corriente\s*:", a["descripcion"]):
                    comp = "inversion|corriente"
                    partes = re.split(r"Corriente\s*:", a["descripcion"], maxsplit=1)
                    t1, an1, m1, n1 = parse_metas(partes[0], anio, a["fecha_fin"])
                    t2, an2, m2, n2 = parse_metas(partes[1], anio, a["fecha_fin"])
                    for cmp_, t_, m_ in (("inversion", t1, m1), ("corriente", t2, m2)):
                        for mes, v in sorted(m_.items()):
                            metas.append({"accion_id": None, "componente": cmp_, "tipo_meta": t_, "periodo": f"{anio}-{mes:02d}", "meta_mes": fnum(v)})
                    tipo_meta, anual, nota = f"{t1}|{t2}", f"{fnum(an1)}|{fnum(an2)}", f"Dos componentes. {n1} / {n2}"
                    metas_mes = {}
                match = {
                    "accion_id": slug(a["titulo"], usados), "institucion": "MTDH", "titulo": a["titulo"],
                    "descripcion": a["descripcion"], "tipo": a["tipo"] or "", "financiamiento": a["financiamiento"] or "",
                    "fuente_financiamiento": a["fuente"] or "", "frecuencia": a["frecuencia"] or "",
                    "cobertura_territorial": a["cobertura"] or "", "fecha_inicio": a["fecha_inicio"], "fecha_fin": a["fecha_fin"],
                    "tipo_meta": tipo_meta, "meta_anual": anual if isinstance(anual, str) else fnum(anual),
                    "componentes": comp, "corriente_planificado": fnum(a["corriente"]), "inversion_planificado": fnum(a["inversion"]),
                    "nota_meta": nota,
                }
                catalogo.append(match)
                for m in metas:
                    if m["accion_id"] is None:
                        m["accion_id"] = match["accion_id"]
                for mes, v in sorted(metas_mes.items()):
                    metas.append({"accion_id": match["accion_id"], "componente": "", "tipo_meta": tipo_meta,
                                  "periodo": f"{anio}-{mes:02d}", "meta_mes": fnum(v)})
                for ac in [x for x in f["actores"] if x["num"] == a["num"]]:
                    actores.append({"accion_id": match["accion_id"], **{k: ac[k] for k in ("actor", "rol", "descripcion_rol", "incidencia")}})
            ya_principal = any(p["accion_id"] == match["accion_id"] for p in puente)
            puente.append({"indicador_id": iid, "accion_id": match["accion_id"], "num_en_ficha": a["num"],
                           "es_principal": "0" if ya_principal else "1"})

    escribir_csv(destino, catalogo, list(catalogo[0].keys()))
    escribir_csv(os.path.join(dir_anio, "indicador_accion.csv"), puente, ["indicador_id", "accion_id", "num_en_ficha", "es_principal"])
    escribir_csv(os.path.join(dir_anio, "metas_programadas.csv"), metas, ["accion_id", "componente", "tipo_meta", "periodo", "meta_mes"])
    escribir_csv(os.path.join(dir_anio, "actores.csv"), actores, ["accion_id", "actor", "rol", "descripcion_rol", "incidencia"])
    for n in ("seguimiento.csv", "justificaciones.csv"):
        ruta = os.path.join(dir_anio, n)
        if not os.path.exists(ruta):
            escribir_csv(ruta, [], CAMPOS_SEG if n == "seguimiento.csv" else CAMPOS_JUST)
    print(f"Catálogo {anio}: {len(catalogo)} acciones únicas, {len(puente)} vínculos indicador-acción, "
          f"{len(metas)} metas mensuales, {len(actores)} actores.")
    sin_meta = [c["accion_id"] for c in catalogo if "SIN META" in c["nota_meta"]]
    if sin_meta:
        print("  ! Revisar metas de:", ", ".join(sin_meta))
    for n, motivo in ignoradas:
        print(f"  - Ignorada: {n} -> {motivo}")
    print("BORRADOR generado. Revise acciones.csv y metas_programadas.csv antes de la primera carga mensual.")


# ----------------------------------------------------------------------------- ACTUALIZAR (mensual)

CAMPOS_SEG = ["anio", "periodo", "accion_id", "componente", "resultado", "unidad_resultado", "tipo_meta",
              "meta_anual", "meta_a_la_fecha", "avance_anual_pct", "cumplimiento_fecha_pct", "avance_reportado_pct",
              "estado", "corte_reportado", "corr_codificado", "corr_reprogramado", "corr_devengado",
              "inv_codificado", "inv_reprogramado", "inv_devengado", "presupuesto_vigente", "presupuesto_devengado",
              "ejecucion_presupuestaria_pct", "fichas_fuente", "validacion", "fecha_carga"]
CAMPOS_JUST = ["anio", "periodo", "accion_id", "componente", "justificacion", "enlaces", "dificultades"]


def corte_desde_texto(txt, anio):
    t = norm(txt)
    pats = [rf"enero\s*(?:a|y|hasta|al)?\s*(?:el mes de\s*)?({'|'.join(MES_NUM)})\s*(?:de|del)?\s*(?:{anio})?",
            rf"(?:al mes de|en el mes de|mes de)\s*({'|'.join(MES_NUM)})",
            rf"periodo de\s*({'|'.join(MES_NUM)})"]
    for p in pats:
        m = re.search(p, t)
        if m:
            return f"{anio}-{MES_NUM[m.group(1)]:02d}"
    if "primer semestre" in t:
        return f"{anio}-06"
    return ""


def extraer_componente(raw, comp):
    """'Inversión: 10.963\\n\\nCorriente: 7918' -> valor del componente."""
    s = str(raw or "")
    etiqueta = {"inversion": r"Inversi[oó]n", "corriente": r"Corriente"}[comp]
    m = re.search(etiqueta + r"\s*:\s*([\d.,]+\s*%?)", s, re.I)
    return m.group(1) if m else None


def a_porcentaje(v):
    """Avance reportado: 0.5939 -> 59.39 ; '69%' -> 69 ; 1.2313 -> 123.13"""
    if v is None:
        return None
    if isinstance(v, str) and "%" in v:
        return parse_num(v)
    x = parse_num(v)
    return None if x is None else x * 100


def cmd_actualizar(args):
    periodo = args.periodo
    if not re.fullmatch(r"\d{4}-\d{2}", periodo):
        sys.exit("--periodo debe tener forma AAAA-MM")
    anio, mes = int(periodo[:4]), int(periodo[5:])
    dir_anio = os.path.join(args.datos, str(anio))
    indicadores = sorted(leer_csv(os.path.join(args.datos, args.indicadores)), key=lambda r: int(r["orden"]))
    catalogo = {r["accion_id"]: r for r in leer_csv(os.path.join(dir_anio, "acciones.csv"))}
    if not catalogo:
        sys.exit(f"No hay catálogo para {anio}. Ejecute primero: inicializar --anio {anio} --fichas ...")
    puente = leer_csv(os.path.join(dir_anio, "indicador_accion.csv"))
    por_num = {(p["indicador_id"], int(p["num_en_ficha"])): p for p in puente}
    metas = defaultdict(dict)
    tipo_por = {}
    for m in leer_csv(os.path.join(dir_anio, "metas_programadas.csv")):
        k = (m["accion_id"], m["componente"])
        metas[k][int(m["periodo"][5:7])] = parse_num(m["meta_mes"])
        tipo_por[k] = m["tipo_meta"]
    seg_prev_all = leer_csv(os.path.join(dir_anio, "seguimiento.csv"))
    periodos_prev = sorted({r["periodo"] for r in seg_prev_all if r["periodo"] < periodo})
    prev = {(r["accion_id"], r["componente"]): r for r in seg_prev_all if periodos_prev and r["periodo"] == periodos_prev[-1]}

    fichas, ignoradas = leer_fichas(args.fichas, indicadores)
    hallazgos = []  # (severidad, codigo, indicador/acción, detalle)

    def H(sev, cod, donde, det):
        hallazgos.append((sev, cod, donde, det))

    for ind in indicadores:
        if ind["indicador_id"] not in fichas:
            H("ADVERTENCIA", "F1", ind["indicador_id"], "No llegó ficha este mes para este indicador.")
    for iid, f in fichas.items():
        ind = next(i for i in indicadores if i["indicador_id"] == iid)
        txt_meta = norm(f["info"]["meta"])
        for campo in ("linea_base", "meta_2029"):
            v = ind[campo].replace(".", ",")
            if norm(v) not in txt_meta and ind[campo] not in f["info"]["meta"]:
                H("ERROR", "F2", iid, f"La {campo.replace('_', ' ')} ({ind[campo]}) no coincide con el texto de la meta en la ficha: «{f['info']['meta']}».")

    # 1) recoger observaciones crudas por acción/componente
    obs = defaultdict(list)
    for iid, f in fichas.items():
        dif = f["dificultades"]
        for s in f["seg"]:
            p = por_num.get((iid, s["num"]))
            if not p:
                H("ERROR", "E1", f"{iid} · acción {s['num']}", f"Acción no catalogada: «{s['titulo'][:90]}». ¿Acción nueva o reenumerada? Agregar a acciones.csv / indicador_accion.csv.")
                continue
            aid = p["accion_id"]
            cat = catalogo[aid]
            if difflib.SequenceMatcher(None, norm(s["titulo"]), norm(cat["titulo"])).ratio() < 0.8:
                H("ERROR", "E1b", f"{iid} · acción {s['num']}", f"El número {s['num']} ya no corresponde a «{cat['titulo'][:60]}»; la ficha dice «{s['titulo'][:60]}». ¿Se reenumeraron las acciones?")
            comps = cat["componentes"].split("|") if cat["componentes"] else [""]
            for comp in comps:
                raw_r = extraer_componente(s["resultado_raw"], comp) if comp else s["resultado_raw"]
                raw_a = extraer_componente(s["avance_raw"], comp) if comp else s["avance_raw"]
                obs[(aid, comp)].append({"iid": iid, "principal": p["es_principal"] == "1", "archivo": f["archivo"],
                                         "s": s, "raw_r": raw_r, "raw_a": raw_a, "dificultad": dif.get(s["num"], "")})

    # 2) consolidar
    nuevas, justs = [], []
    ahora = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    for (aid, comp), lista in obs.items():
        cat = catalogo[aid]
        donde = f"{aid}{' · ' + comp if comp else ''}"
        base = next((o for o in lista if o["principal"]), lista[0])
        s = base["s"]
        cods = []
        # duplicados en varias fichas
        if len(lista) > 1:
            for o in lista:
                if o is base:
                    continue
                difs = []
                if parse_num(o["raw_r"]) != parse_num(base["raw_r"]):
                    difs.append(f"resultado {o['raw_r']} vs {base['raw_r']}")
                ra, rb = a_porcentaje(o["raw_a"]), a_porcentaje(base["raw_a"])
                if ra is not None and rb is not None and abs(ra - rb) > 0.1:
                    difs.append(f"avance {ra:.1f} % vs {rb:.1f} %")
                for c in ("corr_devengado", "inv_devengado", "corr_codificado", "inv_codificado", "corr_reprogramado", "inv_reprogramado"):
                    x, y = o["s"][c] or 0, s[c] or 0
                    if abs(x - y) > 1:
                        difs.append(f"{c} {x:,.2f} vs {y:,.2f}")
                if o["s"]["estado"] != s["estado"]:
                    difs.append(f"estado «{o['s']['estado']}» vs «{s['estado']}»")
                if difs:
                    H("ERROR", "E2", donde, f"La misma acción trae datos distintos en {o['archivo']} y {base['archivo']}: " + "; ".join(difs)
                      + ". Se publica la ficha principal y el avance se recalcula con la meta.")
                    cods.append("E2")
        # resultado
        tipo = tipo_por.get((aid, comp)) or (cat["tipo_meta"].split("|")[0] if cat["tipo_meta"] else "anual")
        res = parse_num(base["raw_r"])
        unidad = ""
        if tipo in ("hito", "tasa") and res is not None and res <= 1.5:
            res, unidad = res * 100, "%"
        if res is None:
            H("ERROR", "E5", donde, f"Resultado no numérico: «{str(base['raw_r'])[:80]}».")
            cods.append("E5")
        # metas
        mm = metas.get((aid, comp), {})
        meta_anual = None
        if mm:
            meta_anual = (parse_num(cat["meta_anual"]) if tipo == "tasa" else
                          max(mm.values()) if tipo in ("cobertura", "hito") else sum(v for v in mm.values() if v))
        elif cat["meta_anual"] and "|" not in cat["meta_anual"]:
            meta_anual = parse_num(cat["meta_anual"])
        # si la justificación declara otro corte (p. ej. dato a agosto dentro de la carga de julio),
        # la meta a la fecha se calcula al corte REAL del dato, no al de la carga
        corte = corte_desde_texto(s["justificacion"], anio)
        mes_ref = int(corte[5:]) if corte else mes
        meta_fecha = meta_a_la_fecha(tipo, mm, mes_ref)
        avance_anual = (res / meta_anual * 100) if (res is not None and meta_anual) else None
        cumpl = (res / meta_fecha * 100) if (res is not None and meta_fecha) else None
        rep = a_porcentaje(base["raw_a"])
        if rep is not None and avance_anual is not None and abs(rep - avance_anual) > 0.5:
            alt = f" (coincide con el cumplimiento a la fecha: {cumpl:.1f} %)" if cumpl is not None and abs(rep - cumpl) <= 0.5 else ""
            H("ERROR" if not alt else "INFO", "E3" if not alt else "I3", donde,
              f"Avance reportado {rep:.2f} % ≠ recalculado {avance_anual:.2f} % (resultado {fnum(res)} / meta anual {fnum(meta_anual)}){alt}.")
            cods.append("E3" if not alt else "I3")
        if avance_anual is not None and avance_anual > 100 and tipo != "tasa":
            H("ADVERTENCIA", "A2", donde, f"Avance anual de {avance_anual:.1f} % (supera la meta del año). Confirmar meta o justificar.")
            cods.append("A2")
        if tipo == "anual":
            H("INFO", "I1", donde, "Meta sin programación mensual: no se calcula cumplimiento a la fecha.")
        # corte
        if corte and corte != periodo:
            H("ADVERTENCIA", "A1", donde, f"La justificación reporta corte {corte} y la carga es {periodo} "
              f"(la meta a la fecha se calculó a {corte}).")
            cods.append("A1")
        # fechas y estado
        if s["fecha_fin"] and s["fecha_fin"][:7] < periodo and norm(s["estado"]) not in ("cumplido", "finalizado", "ejecutado", "concluido"):
            H("ADVERTENCIA", "A3", donde, f"Fecha fin {s['fecha_fin']} ya pasó y el estado es «{s['estado']}».")
            cods.append("A3")
        # presupuesto (solo en el componente principal para no duplicar)
        pres = {k: s[k] for k in ("corr_codificado", "corr_reprogramado", "corr_devengado", "inv_codificado", "inv_reprogramado", "inv_devengado")}
        if comp and comp != (cat["componentes"].split("|")[0]):
            pres = {k: None for k in pres}
        vig_c = pres["corr_reprogramado"] or pres["corr_codificado"] or 0
        vig_i = pres["inv_reprogramado"] or pres["inv_codificado"] or 0
        dev = (pres["corr_devengado"] or 0) + (pres["inv_devengado"] or 0)
        for tag, cod, rep_, devv in (("corriente", pres["corr_codificado"], pres["corr_reprogramado"], pres["corr_devengado"]),
                                     ("inversión", pres["inv_codificado"], pres["inv_reprogramado"], pres["inv_devengado"])):
            if (devv or 0) > 0 and not cod:
                H("ERROR", "E4", donde, f"Presupuesto {tag}: codificado vacío/0 con devengado {devv:,.2f}" + (f" (reprogramado {rep_:,.2f})" if rep_ else "") + ".")
                cods.append("E4")
            if rep_ and not cod and not devv:
                H("ADVERTENCIA", "A5", donde, f"Presupuesto {tag}: hay {rep_:,.2f} en «reprogramación» pero la acción no tiene codificado "
                  "ni devengado (¿valor en la columna equivocada?).")
                cods.append("A5")
            vig = rep_ or cod or 0
            if vig and (devv or 0) > vig + 1:
                H("ADVERTENCIA", "A4", donde, f"Presupuesto {tag}: devengado {devv:,.2f} supera el vigente {vig:,.2f}.")
                cods.append("A4")
        planificado = (parse_num(cat["corriente_planificado"]) or 0) + (parse_num(cat["inversion_planificado"]) or 0)
        if planificado and (vig_c + vig_i) and abs((vig_c + vig_i) - planificado) > 1 and not comp[:1] == "c":
            H("INFO", "I2", donde, f"Presupuesto vigente {vig_c + vig_i:,.2f} ≠ planificado {planificado:,.2f} (reprogramación).")
        # enlaces
        if not enlaces(s["justificacion"]):
            H("ADVERTENCIA", "A6", donde, "La justificación no trae enlace verificable.")
            cods.append("A6")
        # regresión vs corte anterior
        p_ = prev.get((aid, comp))
        if p_ and res is not None and parse_num(p_["resultado"]) is not None and tipo in ("flujo", "trimestral", "semestral", "hito"):
            if res < parse_num(p_["resultado"]):
                H("ADVERTENCIA", "A7", donde, f"El resultado acumulado bajó: {p_['resultado']} ({p_['periodo']}) → {fnum(res)} ({periodo}).")
                cods.append("A7")
        vig_total = vig_c + vig_i
        nuevas.append({
            "anio": anio, "periodo": periodo, "accion_id": aid, "componente": comp,
            "resultado": fnum(res), "unidad_resultado": unidad, "tipo_meta": tipo,
            "meta_anual": fnum(meta_anual), "meta_a_la_fecha": fnum(meta_fecha),
            "avance_anual_pct": fnum(avance_anual), "cumplimiento_fecha_pct": fnum(cumpl), "avance_reportado_pct": fnum(rep),
            "estado": s["estado"], "corte_reportado": corte,
            **{k: fnum(v) for k, v in pres.items()},
            "presupuesto_vigente": fnum(vig_total) if vig_total else "", "presupuesto_devengado": fnum(dev) if (vig_total or dev) else "",
            "ejecucion_presupuestaria_pct": fnum(dev / vig_total * 100) if vig_total else "",
            "fichas_fuente": " | ".join(sorted({o["archivo"] for o in lista})),
            "validacion": ";".join(sorted(set(cods))), "fecha_carga": ahora,
        })
        justs.append({"anio": anio, "periodo": periodo, "accion_id": aid, "componente": comp,
                      "justificacion": s["justificacion"], "enlaces": " | ".join(enlaces(s["justificacion"])),
                      "dificultades": base["dificultad"]})

    # acciones del catálogo que no llegaron
    llegaron = {k[0] for k in obs}
    for aid in catalogo:
        if aid not in llegaron:
            H("ADVERTENCIA", "F3", aid, "Acción del catálogo sin reporte en las fichas de este mes.")

    # 3) escribir (idempotente por periodo)
    orden_acc = {a: i for i, a in enumerate(catalogo)}
    seg = [r for r in seg_prev_all if r["periodo"] != periodo] + nuevas
    seg.sort(key=lambda r: (r["periodo"], orden_acc.get(r["accion_id"], 999), r["componente"]))
    escribir_csv(os.path.join(dir_anio, "seguimiento.csv"), seg, CAMPOS_SEG)
    js = [r for r in leer_csv(os.path.join(dir_anio, "justificaciones.csv")) if r["periodo"] != periodo] + justs
    js.sort(key=lambda r: (r["periodo"], orden_acc.get(r["accion_id"], 999), r["componente"]))
    escribir_csv(os.path.join(dir_anio, "justificaciones.csv"), js, CAMPOS_JUST)

    # 4) reporte
    ruta_rep = escribir_reporte(args, periodo, fichas, ignoradas, hallazgos, nuevas, prev, periodos_prev, catalogo)
    n = defaultdict(int)
    for h in hallazgos:
        n[h[0]] += 1
    print(f"Periodo {periodo}: {len(fichas)} fichas, {len(nuevas)} filas de seguimiento "
          f"({len(catalogo)} acciones en catálogo). Errores: {n['ERROR']} · Advertencias: {n['ADVERTENCIA']} · Info: {n['INFO']}")
    print(f"Reporte: {ruta_rep}")
    return 1 if (args.estricto and n["ERROR"]) else 0


def escribir_reporte(args, periodo, fichas, ignoradas, hallazgos, nuevas, prev, periodos_prev, catalogo):
    os.makedirs(args.reportes, exist_ok=True)
    ruta = os.path.join(args.reportes, f"validacion_{periodo}.md")
    orden = {"ERROR": 0, "ADVERTENCIA": 1, "INFO": 2}
    hallazgos = sorted(hallazgos, key=lambda h: (orden[h[0]], h[1], h[2]))
    n = defaultdict(int)
    for h in hallazgos:
        n[h[0]] += 1
    L = [f"# Validación PND · Protección Social — corte {periodo}", "",
         f"Generado: {dt.datetime.now():%Y-%m-%d %H:%M} · Fichas leídas: {len(fichas)} · Filas cargadas: {len(nuevas)}", "",
         f"**Errores: {n['ERROR']}** (revisar con Planificación antes de publicar) · Advertencias: {n['ADVERTENCIA']} · Informativos: {n['INFO']}", ""]
    L += ["## Fichas", "", "| Indicador | Archivo |", "|---|---|"]
    L += [f"| {iid} | {f['archivo']} |" for iid, f in fichas.items()]
    for nom, mot in ignoradas:
        L.append(f"| *(ignorada)* | {nom} — {mot} |")
    L += ["", "## Hallazgos", ""]
    if hallazgos:
        L += ["| Severidad | Código | Acción / indicador | Detalle |", "|---|---|---|---|"]
        L += [f"| {s} | {c} | {d} | {t.replace('|', '/')} |" for s, c, d, t in hallazgos]
    else:
        L.append("Sin hallazgos.")
    L += ["", f"## Resumen por acción ({periodo})", "",
          "| Acción | Resultado | Meta a la fecha | Cumplimiento a la fecha | Avance anual | Ejec. presupuestaria | Validación |",
          "|---|---:|---:|---:|---:|---:|---|"]
    for r in nuevas:
        nombre = catalogo[r["accion_id"]]["titulo"][:70] + (f" ({r['componente']})" if r["componente"] else "")
        L.append(f"| {nombre} | {r['resultado']}{r['unidad_resultado']} | {r['meta_a_la_fecha'] or '—'} | "
                 f"{(r['cumplimiento_fecha_pct'] + ' %') if r['cumplimiento_fecha_pct'] else '—'} | "
                 f"{(r['avance_anual_pct'] + ' %') if r['avance_anual_pct'] else '—'} | "
                 f"{(r['ejecucion_presupuestaria_pct'] + ' %') if r['ejecucion_presupuestaria_pct'] else '—'} | {r['validacion'] or 'OK'} |")
    if prev:
        resueltos = []
        for (aid, comp), p in prev.items():
            antes = {c for c in p["validacion"].split(";") if c[:1] in ("E", "A")}
            ahora = next((set(r["validacion"].split(";")) for r in nuevas if (r["accion_id"], r["componente"]) == (aid, comp)), None)
            if ahora is not None and antes - ahora:
                resueltos.append((aid + (f" · {comp}" if comp else ""), ", ".join(sorted(antes - ahora))))
        L += ["", f"## Hallazgos resueltos respecto a {periodos_prev[-1]}", ""]
        L += (["| Acción | Códigos que ya no aparecen |", "|---|---|"] + [f"| {a} | {c} |" for a, c in resueltos]) if resueltos else ["Ninguno."]
        L += ["", f"## Cambios frente al corte anterior ({periodos_prev[-1]} → {periodo})", "",
              "| Acción | Resultado antes | Resultado ahora | Δ | Avance anual antes | Avance anual ahora |", "|---|---:|---:|---:|---:|---:|"]
        for r in nuevas:
            p = prev.get((r["accion_id"], r["componente"]))
            if not p:
                continue
            a, b = parse_num(p["resultado"]), parse_num(r["resultado"])
            delta = fnum(b - a) if (a is not None and b is not None) else ""
            nombre = catalogo[r["accion_id"]]["titulo"][:60] + (f" ({r['componente']})" if r["componente"] else "")
            L.append(f"| {nombre} | {p['resultado']} | {r['resultado']} | {delta} | {p['avance_anual_pct'] or '—'} | {r['avance_anual_pct'] or '—'} |")
    L += ["", "### Códigos", "",
          "- **E1/E1b** acción no catalogada o reenumerada · **E2** misma acción con datos distintos entre fichas · "
          "**E3** avance reportado no reproducible con la meta · **E4** codificado vacío con devengado · **E5** resultado no numérico · "
          "**F2** línea base/meta 2029 distinta a la registrada",
          "- **A1** corte de la justificación distinto al periodo · **A2** avance > 100 % · **A3** fecha fin vencida sin cierre · "
          "**A4** devengado > vigente · **A5** reprogramado sin codificado ni devengado · **A6** sin enlace verificable · **A7** acumulado bajó · **F1/F3** falta ficha / acción sin reporte",
          "- **I1** meta anual sin programación · **I2** reprogramación presupuestaria · **I3** el % reportado corresponde al cumplimiento a la fecha, no al anual"]
    with open(ruta, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L) + "\n")
    return ruta


# ----------------------------------------------------------------------------- CLI


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("inicializar", help="Genera el borrador del catálogo anual desde la matriz de planificación")
    a.add_argument("--anio", required=True)
    a.add_argument("--fichas", required=True)
    a.add_argument("--forzar", action="store_true")
    b = sub.add_parser("actualizar", help="Carga el seguimiento de un mes")
    b.add_argument("--periodo", required=True, help="AAAA-MM")
    b.add_argument("--fichas", required=True)
    b.add_argument("--estricto", action="store_true", help="Termina con código 1 si hay errores")
    for p in (a, b):
        p.add_argument("--datos", default=DATOS_DEFAULT)
        p.add_argument("--reportes", default=REPORTES_DEFAULT)
        p.add_argument("--indicadores", default="pnd_ps_indicadores.csv",
                       help="Tabla de indicadores dentro de --datos (Trabajo: --datos datos/pnd/trabajo --indicadores indicadores.csv)")
    args = ap.parse_args()
    sys.exit(cmd_inicializar(args) if args.cmd == "inicializar" else cmd_actualizar(args))


if __name__ == "__main__":
    main()
