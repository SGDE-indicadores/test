#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
actualizar_jubilados.py
=======================
Carga mensual del KPI «Número de exservidores y extrabajadores jubilados compensados
económicamente» (KPIs · Trabajo) en el portal.

Lee el Excel que entrega el área y escribe:
  datos/kpis/jubilados/gestion.csv     tabla larga: anio × dimensión (total, sexo, régimen, provincia)
  datos/kpis/jubilados/mensual.csv     serie mensual (solo con el formato nuevo, opciones A o B)
  datos/kpis/jubilados/corte.csv       metadatos del corte cargado
  datos/kpis/kpi_trabajo.csv                tarjeta del «Panorama destacado» (solo el registro 1)
  reportes/kpis/validacion_jubilados_AAAA-MM.md

Acepta dos formatos de entrada (se detectan solos):
  1. Formato actual  (GESTION_POR_GENERO_PROVINCIA_REGIMEN.xlsx): hojas POR PROVINCIA, POR REGIMEN,
     GENERO, con bloques por año uno al lado del otro (categoría, EXP., MONTO).
  2. Formato de entrega mensual (Formato_entrega_jubilados_MTDH.xlsx): hoja «Expedientes»
     (opción A, una fila por expediente) o «Agregado» (opción B). Solo cuenta estado = Pagado.

El archivo trae siempre el histórico completo: cada corte reemplaza la carga anterior
(idempotente). El año del corte se marca como parcial (enero–mes de corte).

Uso:
  python scripts/actualizar_jubilados.py --archivo "C:/jubilados/JUBILADOS_2026-08.xlsx" --corte 2026-08
"""
import argparse
import csv
import os
import re
import sys
import unicodedata
from collections import defaultdict
from datetime import date, datetime

import openpyxl

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]
PROVINCIAS = [("01", "Azuay"), ("02", "Bolívar"), ("03", "Cañar"), ("04", "Carchi"), ("05", "Cotopaxi"),
              ("06", "Chimborazo"), ("07", "El Oro"), ("08", "Esmeraldas"), ("09", "Guayas"), ("10", "Imbabura"),
              ("11", "Loja"), ("12", "Los Ríos"), ("13", "Manabí"), ("14", "Morona Santiago"), ("15", "Napo"),
              ("16", "Pastaza"), ("17", "Pichincha"), ("18", "Tungurahua"), ("19", "Zamora Chinchipe"),
              ("20", "Galápagos"), ("21", "Sucumbíos"), ("22", "Orellana"), ("23", "Santo Domingo de los Tsáchilas"),
              ("24", "Santa Elena"), ("90", "Zonas no delimitadas")]
REGIMENES = {"CODIGO DEL TRABAJO": "Código del Trabajo", "LOEI": "LOEI", "LOSEP": "LOSEP", "LOSE": "LOSE"}
SEXOS = {"FEMENINO": "Mujeres", "MUJER": "Mujeres", "MASCULINO": "Hombres", "HOMBRE": "Hombres"}
CAMPOS = ["anio", "mes_inicio", "mes_fin", "parcial", "dimension", "categoria", "codigo", "expedientes", "monto_usd"]


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().upper()


PROV_IDX = {norm(n): (c, n) for c, n in PROVINCIAS}
PROV_IDX.update({c: (c, n) for c, n in PROVINCIAS})
PROV_IDX["SANTO DOMINGO"] = ("23", "Santo Domingo de los Tsáchilas")


class Log:
    def __init__(self):
        self.e, self.a, self.i = [], [], []


def categoria(dim, valor, log, donde):
    k = norm(valor)
    if dim == "provincia":
        if k not in PROV_IDX:
            log.e.append(f"{donde}: provincia no reconocida «{valor}».")
            return None, str(valor).strip()
        return PROV_IDX[k]
    if dim == "regimen":
        if k not in REGIMENES:
            log.a.append(f"{donde}: régimen fuera de catálogo «{valor}» (se publica tal cual).")
            return "", str(valor).strip()
        if k == "LOSE":
            log.a.append(f"{donde}: régimen «LOSE» pendiente de confirmar (¿Ley Orgánica del Servicio Exterior o LOSEP?).")
        if str(valor) != str(valor).strip():
            log.i.append(f"{donde}: «{valor}» traía espacios al final; se normalizó.")
        return "", REGIMENES[k]
    if dim == "sexo":
        if k not in SEXOS:
            log.e.append(f"{donde}: sexo no reconocido «{valor}».")
            return "", str(valor).strip()
        return "", SEXOS[k]
    return "", str(valor).strip()


def num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


# ------------------------------------------------------------------ formato actual (bloques por año)
def leer_bloques(wb, log):
    dims = {}
    for ws in wb.worksheets:
        n = norm(ws.title)
        dim = "provincia" if "PROVINCIA" in n else "regimen" if "REGIMEN" in n else "sexo" if ("GENERO" in n or "SEXO" in n) else None
        if not dim:
            log.i.append(f"Hoja «{ws.title}» ignorada (no es provincia, régimen ni género).")
            continue
        anios = [(c, int(ws.cell(1, c).value)) for c in range(1, ws.max_column + 1)
                 if isinstance(ws.cell(1, c).value, (int, float)) and 2000 < ws.cell(1, c).value < 2100]
        for c0, anio in anios:
            for r in range(3, ws.max_row + 1):
                k, e, m = ws.cell(r, c0).value, ws.cell(r, c0 + 1).value, ws.cell(r, c0 + 2).value
                if k is None or norm(k).startswith("TOTAL"):
                    continue
                donde = f"{ws.title} {anio} fila {r}"
                e, m = num(e), num(m)
                if e is None or m is None:
                    log.e.append(f"{donde}: expedientes o monto no numéricos ({ws.cell(r, c0 + 1).value!r}, {ws.cell(r, c0 + 2).value!r}).")
                    continue
                if e < 0 or m < 0:
                    log.e.append(f"{donde}: valor negativo.")
                if e != int(e):
                    log.e.append(f"{donde}: expedientes no enteros ({e}).")
                if abs(m * 100 - round(m * 100)) > 1e-4:
                    log.i.append(f"{donde} ({k}): monto con fracciones de centavo ({m}); se redondea a 2 decimales.")
                cod, cat = categoria(dim, k, log, donde)
                clave = (anio, dim, cat)
                acc = dims.setdefault(clave, [cod, 0, 0.0])
                acc[1] += int(e)
                acc[2] += m
    return dims


# ------------------------------------------------------------------ formato nuevo (opciones A y B)
def a_fecha(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return datetime.strptime(str(v)[:10], "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


def leer_plantilla(wb, corte, log):
    """Devuelve lista de registros (anio, mes, sexo, regimen, provincia, expedientes, monto) solo de pagados."""
    regs = []
    ya, ma = int(corte[:4]), int(corte[5:])
    if "Expedientes" in wb.sheetnames:
        ws = wb["Expedientes"]
        h = [c.value for c in ws[2]]
        ix = {k: h.index(k) for k in h if k}
        vistos = set()
        for r, fila in enumerate(ws.iter_rows(min_row=3, values_only=True), 3):
            if not fila or fila[ix["id_expediente"]] in (None, ""):
                continue
            idx = str(fila[ix["id_expediente"]])
            if idx.upper().startswith("EJEMPLO"):
                continue
            if idx in vistos:
                log.e.append(f"Expedientes fila {r}: id_expediente repetido ({idx}).")
            vistos.add(idx)
            if norm(fila[ix["estado"]]) != "PAGADO":
                continue
            f = a_fecha(fila[ix["fecha_pago"]])
            if not f:
                log.e.append(f"Expedientes fila {r}: estado Pagado sin fecha_pago válida.")
                continue
            if (f.year, f.month) > (ya, ma):
                log.a.append(f"Expedientes fila {r}: fecha_pago {f} posterior al corte; no se cuenta.")
                continue
            regs.append((f.year, f.month, fila[ix["sexo"]], fila[ix["regimen"]], fila[ix["provincia_codigo"]],
                         1, num(fila[ix["monto_pagado_usd"]]) or 0.0, f"Expedientes fila {r}"))
        return regs, "A (expedientes)"
    ws = wb["Agregado"]
    h = [c.value for c in ws[2]]
    ix = {k: h.index(k) for k in h if k}
    for r, fila in enumerate(ws.iter_rows(min_row=3, values_only=True), 3):
        if not fila or fila[ix["mes_pago"]] in (None, ""):
            continue
        mp = str(fila[ix["mes_pago"]])[:7]
        if not re.match(r"^\d{4}-\d{2}$", mp):
            log.e.append(f"Agregado fila {r}: mes_pago «{mp}» no tiene formato AAAA-MM.")
            continue
        if r == 3 and ws.cell(3, 1).fill.fgColor.rgb in ("00EDEDED", "FFEDEDED"):
            continue  # fila gris de ejemplo
        if norm(fila[ix["estado"]]) != "PAGADO":
            continue
        y, m = int(mp[:4]), int(mp[5:])
        if (y, m) > (ya, ma):
            log.a.append(f"Agregado fila {r}: mes {mp} posterior al corte; no se cuenta.")
            continue
        regs.append((y, m, fila[ix["sexo"]], fila[ix["regimen"]], fila[ix["provincia_codigo"]],
                     int(num(fila[ix["expedientes"]]) or 0), num(fila[ix["monto_pagado_usd"]]) or 0.0, f"Agregado fila {r}"))
    return regs, "B (agregado)"


def agregar_plantilla(regs, log):
    dims, mensual = {}, defaultdict(lambda: [0, 0.0])
    for y, m, sx, rg, pv, e, monto, donde in regs:
        mensual[(y, m)][0] += e
        mensual[(y, m)][1] += monto
        for dim, val in (("sexo", sx), ("regimen", rg), ("provincia", str(pv or "").zfill(2) if pv not in (None, "") else "")):
            cod, cat = categoria(dim, val, log, donde)
            acc = dims.setdefault((y, dim, cat), [cod, 0, 0.0])
            acc[1] += e
            acc[2] += monto
    return dims, mensual


# ------------------------------------------------------------------ principal
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--archivo", required=True, help="Excel entregado por el área")
    ap.add_argument("--corte", required=True, help="Mes de corte AAAA-MM (ej. 2026-08)")
    ap.add_argument("--datos", default=os.path.join(REPO, "datos"))
    ap.add_argument("--reportes", default=os.path.join(REPO, "reportes", "kpis"))
    args = ap.parse_args()
    if not re.match(r"^\d{4}-(0[1-9]|1[0-2])$", args.corte):
        sys.exit("--corte debe tener formato AAAA-MM")
    ya, ma = int(args.corte[:4]), int(args.corte[5:])

    log = Log()
    wb = openpyxl.load_workbook(args.archivo, data_only=True)
    mensual = {}
    if "Expedientes" in wb.sheetnames or "Agregado" in wb.sheetnames:
        regs, formato = leer_plantilla(wb, args.corte, log)
        dims, mensual = agregar_plantilla(regs, log)
    else:
        formato = "actual (bloques por año)"
        dims = leer_bloques(wb, log)
    if not dims:
        sys.exit("No se encontraron datos en el archivo.")

    anios = sorted({k[0] for k in dims})
    if max(anios) > ya:
        log.e.append(f"Hay datos de {max(anios)}, posteriores al corte {args.corte}.")

    # ---- consistencia: los totales de cada dimensión deben coincidir en cada año
    filas, totales = [], {}
    for anio in anios:
        tot_dim = {}
        for dim in ("sexo", "regimen", "provincia"):
            sel = [(k[2], v) for k, v in dims.items() if k[0] == anio and k[1] == dim]
            if not sel:
                log.a.append(f"{anio}: no hay desagregación por {dim}.")
                continue
            tot_dim[dim] = (sum(v[1] for _, v in sel), sum(v[2] for _, v in sel))
        if not tot_dim:
            continue
        ref = next(iter(tot_dim.values()))
        for dim, (e, m) in tot_dim.items():
            if e != ref[0] or abs(m - ref[1]) > 1.0:
                log.e.append(f"{anio}: el total por {dim} ({e} exp., ${m:,.2f}) no coincide con el de "
                             f"{next(iter(tot_dim))} ({ref[0]} exp., ${ref[1]:,.2f}).")
            elif abs(m - ref[1]) > 0.004:
                log.i.append(f"{anio}: el total por {dim} difiere en ${abs(m - ref[1]):.2f} (redondeo).")
        totales[anio] = ref
        parcial = anio == ya and ma < 12
        base = {"anio": anio, "mes_inicio": 1, "mes_fin": ma if anio == ya else 12, "parcial": "TRUE" if parcial else "FALSE"}
        filas.append({**base, "dimension": "total", "categoria": "Total", "codigo": "", "expedientes": ref[0], "monto_usd": f"{ref[1]:.2f}"})
        orden = {"sexo": ["Mujeres", "Hombres"], "regimen": list(REGIMENES.values())}
        for dim in ("sexo", "regimen", "provincia"):
            sel = sorted([(k[2], v) for k, v in dims.items() if k[0] == anio and k[1] == dim],
                         key=lambda x: (orden.get(dim, []).index(x[0]) if x[0] in orden.get(dim, []) else 99, x[1][0] or "", x[0]))
            for cat, (cod, e, m) in sel:
                filas.append({**base, "dimension": dim, "categoria": cat, "codigo": cod, "expedientes": e, "monto_usd": f"{m:.2f}"})
    if ya in anios and min(anios) < ya:
        log.i.append(f"{ya} comprende enero–{MESES[ma - 1]} ({ma} meses): no es comparable directamente con años completos.")
    if min(anios) in totales and len(anios) > 1 and totales[min(anios)][0] < 0.3 * max(v[0] for v in totales.values()):
        log.a.append(f"{min(anios)} tiene muchos menos expedientes que los demás años ({totales[min(anios)][0]}): "
                     "confirmar si es un año parcial o el inicio del proyecto.")

    if log.e:
        print("ERRORES: no se escribió nada.\n  " + "\n  ".join(log.e))
        escribir_reporte(args, formato, anios, totales, log)
        sys.exit(1)

    dir_j = os.path.join(args.datos, "kpis", "jubilados")
    os.makedirs(dir_j, exist_ok=True)
    escribir(os.path.join(dir_j, "gestion.csv"), filas, CAMPOS)
    escribir(os.path.join(dir_j, "mensual.csv"),
             [{"periodo": f"{y}-{m:02d}", "expedientes": v[0], "monto_usd": f"{v[1]:.2f}"} for (y, m), v in sorted(mensual.items())],
             ["periodo", "expedientes", "monto_usd"])
    escribir(os.path.join(dir_j, "corte.csv"),
             [{"corte": args.corte, "archivo": os.path.basename(args.archivo), "formato": formato,
               "fecha_carga": date.today().isoformat(), "anio_inicio": min(anios), "anio_fin": max(anios)}],
             ["corte", "archivo", "formato", "fecha_carga", "anio_inicio", "anio_fin"])
    panorama(args, filas, totales, anios, ya, ma)
    escribir_reporte(args, formato, anios, totales, log)
    print(f"Jubilados: corte {args.corte} · formato {formato} · {len(anios)} años · {len(filas)} filas. "
          f"Errores 0 · Advertencias {len(log.a)} · Info {len(log.i)}")


def escribir(ruta, filas, campos):
    with open(ruta, "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=campos, extrasaction="ignore")
        w.writeheader()
        w.writerows(filas)


def mill(v):
    return f"${v / 1e6:,.1f} M".replace(",", "X").replace(".", ",").replace("X", ".")


def miles(v):
    return f"{int(round(v)):,}".replace(",", ".")


def cuota(x, total):
    return f"{miles(x)} ({x / total * 100:.1f} %)".replace(".", "#").replace("#", ".", miles(x).count(".")).replace("#", ",") if total else ""


def panorama(args, filas, totales, anios, ya, ma):
    """Reemplaza las filas del registro 1 en datos/kpis/kpi_trabajo.csv (carrusel de KPIs · Trabajo)."""
    ruta = os.path.join(args.datos, "kpis", "kpi_trabajo.csv")
    with open(ruta, encoding="utf-8") as fh:
        previas = list(csv.DictReader(fh))
    campos = list(previas[0].keys())
    for extra in ("serie_titulo", "serie_nota"):
        if extra not in campos:
            campos.append(extra)
    otras = [r for r in previas if r.get("registro_id") != "1"]
    ult = max(anios)
    e, m = totales[ult]
    per = f"enero–{MESES[ma - 1]} {ult}" if ult == ya and ma < 12 else str(ult)
    sx = {r["categoria"]: int(r["expedientes"]) for r in filas if r["anio"] == ult and r["dimension"] == "sexo"}
    rg = sorted([r for r in filas if r["anio"] == ult and r["dimension"] == "regimen"], key=lambda r: -int(r["expedientes"]))
    base = {k: "" for k in campos}
    base.update({"area": "kpis-trabajo", "registro_id": "1", "indicador": "Trabajo", "icono": "doc",
                 "color": "var(--azul)", "color_claro": "var(--azul-pale)", "es_dato_real": "TRUE",
                 "nota": f"Dato real: reporte de gestión de compensación jubilar ({os.path.basename(args.archivo)}), corte {args.corte}."})
    pan = [{**base, "tipo_registro": "desagregacion_actual", "headline": miles(e),
            "headline_sufijo": f"expedientes · {per} · {mill(m)}",
            "tendencia_direccion": "", "tendencia_texto": f"Acumulado {min(anios)}–{ult}: {miles(sum(v[0] for v in totales.values()))} expedientes · "
                                                          f"{mill(sum(v[1] for v in totales.values()))}",
            "impacto_texto": "Cumplir con la compensación jubilar de exservidores y extrabajadores del sector público.",
            "tabulado_encabezado_categoria": per[0].upper() + per[1:], "tabulado_encabezado_valor": "Expedientes",
            "desagregacion_categoria": "Mujeres", "desagregacion_valor": cuota(sx.get('Mujeres', 0), e)}]
    pan.append({**base, "tipo_registro": "desagregacion_actual", "desagregacion_categoria": "Hombres",
                "desagregacion_valor": cuota(sx.get('Hombres', 0), e)})
    for r in rg:
        pan.append({**base, "tipo_registro": "desagregacion_actual", "desagregacion_categoria": "Régimen " + r["categoria"],
                    "desagregacion_valor": miles(int(r["expedientes"]))})
    for a in anios:
        pan.append({**base, "tipo_registro": "serie_historica", "anio": str(a) + (" (ene–" + MESES[ma - 1][:3] + ")" if a == ya and ma < 12 else ""),
                    "valor": totales[a][0], "serie_titulo": "Expedientes por año",
                    "serie_nota": f"{ya} comprende enero–{MESES[ma - 1]}; los demás años son completos." if ma < 12 else ""})
    escribir(ruta, pan + otras, campos)


def escribir_reporte(args, formato, anios, totales, log):
    os.makedirs(args.reportes, exist_ok=True)
    ruta = os.path.join(args.reportes, f"validacion_jubilados_{args.corte}.md")
    L = [f"# Validación · KPI Trabajo › Jubilados · corte {args.corte}", "",
         f"- Archivo: `{os.path.basename(args.archivo)}` · formato {formato}",
         f"- Años: {', '.join(map(str, anios))}", "", "| Año | Expedientes | Monto (USD) |", "|---|---:|---:|"]
    L += [f"| {a} | {totales[a][0]:,} | {totales[a][1]:,.2f} |" for a in anios if a in totales]
    for tit, xs in (("Errores (bloquean la carga)", log.e), ("Advertencias (revisar con el área)", log.a), ("Información", log.i)):
        L += ["", f"## {tit}", ""] + ([f"- {x}" for x in dict.fromkeys(xs)] or ["- Ninguno."])
    with open(ruta, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L) + "\n")
    print("Reporte:", ruta)


if __name__ == "__main__":
    main()
